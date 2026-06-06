#!/usr/bin/env python3
"""Generate the APA-style Amazon Beauty Deliverable 2 report."""

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


REPORT_PATH = "Amazon_Beauty_Recommendation_Deliverable2_Report.docx"
SUMMARY_PATH = Path("outputs/amazon_beauty_final_model_summary.csv")
GRID_PATH = Path("outputs/amazon_beauty_grid_search_results.csv")


def load_results():
    if not SUMMARY_PATH.exists() or not GRID_PATH.exists():
        return {
            "raw_ratings": "[run model script first]",
            "modeled_ratings": "[run model script first]",
            "unique_users": "[run model script first]",
            "unique_products": "[run model script first]",
            "train_ratings": "[run model script first]",
            "test_ratings": "[run model script first]",
            "best_n_factors": "[run model script first]",
            "best_n_epochs": "[run model script first]",
            "best_lr_all": "[run model script first]",
            "test_rmse": "[run model script first]",
            "test_mae": "[run model script first]",
            "mean_cv_rmse": "[run model script first]",
        }

    summary = pd.read_csv(SUMMARY_PATH).iloc[0].to_dict()
    grid = pd.read_csv(GRID_PATH)
    best = grid.sort_values("rank_test_rmse").iloc[0]
    summary["mean_cv_rmse"] = round(float(best["mean_test_rmse"]), 4)
    return summary


def fmt(value):
    if isinstance(value, float):
        return f"{value:,.4f}"
    if isinstance(value, int):
        return f"{value:,}"
    return str(value)


def build_sections(results):
    best_n_factors = fmt(results["best_n_factors"])
    best_n_epochs = fmt(results["best_n_epochs"])
    best_lr_all = fmt(results["best_lr_all"])
    cv_rmse = fmt(results["mean_cv_rmse"])
    test_rmse = fmt(results["test_rmse"])
    test_mae = fmt(results["test_mae"])
    modeled_ratings = fmt(results["modeled_ratings"])
    unique_users = fmt(results["unique_users"])
    unique_products = fmt(results["unique_products"])
    train_ratings = fmt(results["train_ratings"])
    test_ratings = fmt(results["test_ratings"])

    return [
        (
            "Hyperparameter Tuning",
            [
                f"The final recommendation model was refined using a reproducible 20,000-rating sample from the Amazon Beauty Product Ratings dataset on Kaggle. The sample includes user identifiers, product identifiers, product categories, rating values, timestamps, and product URLs. The workflow models {modeled_ratings} ratings across {unique_users} users and {unique_products} beauty products. The model uses singular value decomposition (SVD), a matrix factorization method that learns latent user and product factors from rating behavior (Koren et al., 2009).",
                f"Hyperparameter optimization used Surprise GridSearchCV with three-fold cross-validation and the required parameter grid: n_factors = 50 or 100, n_epochs = 20 or 40, and lr_all = 0.002 or 0.005. The best configuration was n_factors = {best_n_factors}, n_epochs = {best_n_epochs}, and lr_all = {best_lr_all}, with a mean cross-validation RMSE of {cv_rmse}. This tuning step improved the model selection process by comparing model complexity, training duration, and learning rate instead of relying on default SVD settings.",
            ],
        ),
        (
            "Final Model Evaluation",
            [
                f"The optimized SVD model was retrained on the training split and evaluated on a held-out testing dataset. The split produced {train_ratings} training ratings and {test_ratings} testing ratings. Final model performance was evaluated with root mean squared error (RMSE) and mean absolute error (MAE). The tuned Amazon Beauty recommender achieved a test RMSE of {test_rmse} and a test MAE of {test_mae}. RMSE is useful because it penalizes large prediction errors, while MAE shows the average absolute distance between predicted product ratings and actual customer ratings.",
                "The exported prediction comparison file supports item-level analysis by placing the actual rating, predicted rating, and absolute error side by side. Smaller errors indicate cases where the model captured broad preference patterns, while larger errors show cases where collaborative filtering was insufficient. These errors can occur because beauty purchases are highly personal and contextual. A customer may rate a product based on skin type, hair texture, allergies, scent preference, packaging quality, shipping condition, price expectations, or brand loyalty, none of which are directly represented in a simple user-product-rating matrix.",
                "The results show that SVD is a reasonable baseline for product recommendation, but the model should not be treated as a production-ready recommendation engine by itself. Offline RMSE and MAE measure rating-prediction accuracy, but they do not fully measure whether recommended products are useful, diverse, fair, explainable, or profitable. A production system should also evaluate ranking metrics such as precision at k, recall at k, normalized discounted cumulative gain, catalog coverage, diversity, novelty, and user satisfaction. Because beauty products often involve personal safety and suitability, recommendation quality should also consider product attributes, customer needs, and post-purchase feedback.",
            ],
        ),
        (
            "Ethical Considerations",
            [
                "Recommendation systems can reinforce bias because they learn from historical behavior. In the Amazon Beauty setting, popularity bias may cause already popular products and brands to receive more exposure while newer, smaller, or niche products receive fewer recommendations. Review bias is also important because ratings may overrepresent customers who are unusually satisfied or dissatisfied, while silent customers are missing from the dataset. Demographic bias can appear indirectly through product categories, price points, geography, language, or beauty standards even when explicit demographic fields are absent. Fair recommender research emphasizes that systems can affect multiple stakeholders, including users, product providers, and platforms (Deldjoo et al., 2024).",
                "Privacy and transparency are also major concerns. Product ratings can reveal sensitive information about grooming habits, health-related needs, gender presentation, age-related concerns, or skin and hair conditions. Users should understand why a product is recommended and should have control over personalization. Fairness research shows that bias can enter through data collection, model objectives, and deployment feedback loops (Pessach & Shmueli, 2022). Mitigation steps include privacy-preserving data practices, opt-out controls, clear explanations, fairness audits across product categories and price tiers, popularity-bias monitoring, and re-ranking strategies that balance relevance with diversity and fair exposure.",
            ],
        ),
        (
            "Real-World Application",
            [
                "This recommendation system could be applied in an e-commerce beauty marketplace to personalize product discovery. The model can estimate how a customer might rate unseen products and then recommend items with the highest predicted ratings. This could support homepage recommendations, product-detail-page suggestions, cart add-ons, email campaigns, and personalized search ranking. For customers, the goal would be faster discovery of relevant beauty products. For the platform, the goal would be higher engagement, conversion, retention, and customer satisfaction.",
                "Several improvements would make the system more effective in production. First, the model should become hybrid by combining collaborative filtering with product metadata such as brand, category, price, ingredients, scent, color, size, skin type, hair type, and review text embeddings. This would reduce cold-start problems for new products and new users. Second, the platform should include ranking and business metrics rather than relying only on RMSE and MAE. Online A/B testing could measure click-through rate, add-to-cart rate, purchase rate, return rate, review sentiment, repeat purchase behavior, and long-term customer satisfaction.",
                "Third, production recommendations should include guardrails for fairness, safety, and diversity. A beauty recommender should avoid repeatedly promoting only dominant brands or narrow beauty standards. It should also avoid unsafe personalization, such as inferring sensitive characteristics without user consent. Re-ranking can balance predicted relevance with product diversity, price diversity, and fair exposure for smaller sellers. Finally, the system should be monitored over time for model drift, changing product catalogs, seasonal demand, manipulated reviews, and unequal recommendation quality across product categories.",
            ],
        ),
        (
            "Final Thoughts and Conclusion",
            [
                f"Deliverable 2 refined an Amazon Beauty product recommendation model by tuning SVD hyperparameters, selecting the best model through GridSearchCV, and evaluating the optimized model on a testing split. The best model used {best_n_factors} latent factors, {best_n_epochs} training epochs, and a learning rate of {best_lr_all}. Its final RMSE of {test_rmse} and MAE of {test_mae} provide a measurable baseline for product rating prediction.",
                "The main conclusion is that product recommendation quality requires more than predictive accuracy. The model must also be evaluated for fairness, transparency, privacy, diversity, and real-world usefulness. Because this is an individual submission, the final review process focused on ensuring that the code runs independently, the methodology is documented, the output files support the evaluation discussion, and the final report follows APA formatting requirements.",
            ],
        ),
    ]


REFERENCES = [
    "Deldjoo, Y., Jannach, D., Bellogin, A., Difonzo, A., & Zanzonelli, D. (2024). Fairness in recommender systems: Research landscape and future directions. User Modeling and User-Adapted Interaction, 34, 59-108. https://doi.org/10.1007/s11257-023-09364-z",
    "Koren, Y., Bell, R., & Volinsky, C. (2009). Matrix factorization techniques for recommender systems. Computer, 42(8), 30-37. https://doi.org/10.1109/MC.2009.263",
    "Pessach, D., & Shmueli, E. (2022). A review on fairness in machine learning. ACM Computing Surveys, 55(3), Article 51. https://doi.org/10.1145/3494672",
    "Ricci, F., Rokach, L., & Shapira, B. (2022). Recommender systems: Techniques, applications, and challenges. In F. Ricci, L. Rokach, & B. Shapira (Eds.), Recommender systems handbook (3rd ed., pp. 1-35). Springer. https://doi.org/10.1007/978-1-0716-2197-4_1",
]


def set_document_style(document):
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 2
    style.paragraph_format.space_after = Pt(0)


def add_paragraph(document, text, first_line_indent=True):
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.line_spacing = 2
    paragraph.paragraph_format.space_after = Pt(0)
    if first_line_indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
    return paragraph


def add_heading(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 2
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def main():
    results = load_results()
    document = Document()
    set_document_style(document)

    for text in [
        "Group Project Part - 2 (Individual Submission)",
        "Mounica Kudumulla",
        "Fundamentals of AI-Enabled Systems (PhDAI-732-A01)",
        "Dr. Marwan Omar",
        "June 6, 2026",
    ]:
        paragraph = add_paragraph(document, text, first_line_indent=False)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_page_break()

    add_heading(document, "Amazon Beauty Product Recommendation System: Deliverable 2")
    add_paragraph(
        document,
        "Dataset: Amazon Beauty Product Ratings from Kaggle",
        first_line_indent=False,
    )

    for heading, paragraphs in build_sections(results):
        add_heading(document, heading)
        for paragraph in paragraphs:
            add_paragraph(document, paragraph)

    add_heading(document, "References")
    for reference in REFERENCES:
        paragraph = add_paragraph(document, reference, first_line_indent=False)
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5)

    document.save(REPORT_PATH)
    print(f"Saved {REPORT_PATH}")


if __name__ == "__main__":
    main()
